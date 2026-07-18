from io import BytesIO

from docx import Document
from fastapi.testclient import TestClient
from reportlab.pdfgen import canvas

from backend.app.main import create_app


def make_docx(text: str) -> bytes:
    stream = BytesIO()
    document = Document()
    document.add_heading("StudyPilot", level=1)
    document.add_paragraph(text)
    document.save(stream)
    return stream.getvalue()


def make_pdf(text: str) -> bytes:
    stream = BytesIO()
    page = canvas.Canvas(stream)
    page.drawString(72, 760, text)
    page.save()
    return stream.getvalue()


def test_imports_real_pdf_docx_markdown_and_text(tmp_path) -> None:
    files = [
        ("rag.pdf", make_pdf("BM25 retrieval baseline"), "application/pdf", "BM25 retrieval baseline"),
        (
            "agent.docx",
            make_docx("Tool calling needs idempotency"),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "Tool calling needs idempotency",
        ),
        ("notes.md", "# RRF\n融合稀疏与稠密召回".encode(), "text/markdown", "融合稀疏与稠密召回"),
        ("plain.txt", "Python 子进程隔离".encode(), "text/plain", "Python 子进程隔离"),
    ]

    with TestClient(create_app(data_dir=tmp_path)) as client:
        imported = []
        for name, content, media_type, expected in files:
            response = client.post("/api/documents/import", files={"file": (name, content, media_type)})
            assert response.status_code == 201, response.text
            assert expected in response.json()["data"]["body"]
            imported.append(response.json()["data"])

    assert len(imported) == 4
    assert all((tmp_path / item["stored_path"]).is_file() for item in imported)


def test_full_text_search_and_highlight_use_document_body(tmp_path) -> None:
    with TestClient(create_app(data_dir=tmp_path)) as client:
        imported = client.post(
            "/api/documents/import",
            files={"file": ("retrieval.md", "# 检索\nBM25 是稀疏检索基线".encode(), "text/markdown")},
        ).json()["data"]
        search = client.get("/api/search", params={"q": "BM25"})
        highlight = client.post(
            f"/api/documents/{imported['id']}/highlights",
            json={"quote": "BM25 是稀疏检索基线", "note": "加入知识点"},
        )

    assert search.status_code == 200
    assert search.json()["data"][0]["document_id"] == imported["id"]
    assert "<mark>BM25</mark>" in search.json()["data"][0]["snippet"]
    assert highlight.status_code == 201


def test_duplicate_document_is_deduplicated(tmp_path) -> None:
    content = "same content".encode()
    with TestClient(create_app(data_dir=tmp_path)) as client:
        first = client.post("/api/documents/import", files={"file": ("a.txt", content, "text/plain")})
        second = client.post("/api/documents/import", files={"file": ("b.txt", content, "text/plain")})

    assert first.status_code == 201
    assert second.status_code == 200
    assert second.json()["data"]["id"] == first.json()["data"]["id"]
    assert second.json()["meta"]["deduplicated"] is True


def test_rejects_unsupported_file_without_copying_it(tmp_path) -> None:
    with TestClient(create_app(data_dir=tmp_path)) as client:
        response = client.post(
            "/api/documents/import", files={"file": ("payload.exe", b"MZ", "application/octet-stream")}
        )

    assert response.status_code == 415
    assert response.json()["error"]["code"] == "UNSUPPORTED_DOCUMENT"
    assert not list((tmp_path / "documents").glob("*"))

