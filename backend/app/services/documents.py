from __future__ import annotations

import hashlib
import json
import re
from io import BytesIO
from pathlib import Path

from docx import Document
from pypdf import PdfReader

from ..db import Database
from ..errors import AppError
from ..repository import as_dict


MAX_DOCUMENT_BYTES = 50 * 1024 * 1024
SUPPORTED = {".pdf", ".docx", ".md", ".txt"}


class DocumentService:
    def __init__(self, database: Database, data_dir: Path) -> None:
        self.database = database
        self.data_dir = data_dir
        self.documents_dir = data_dir / "documents"

    def _course_id(self) -> int:
        with self.database.connect() as connection:
            row = connection.execute(
                "SELECT value_json FROM settings WHERE key = 'active_course'"
            ).fetchone()
        return int(json.loads(row[0]))

    def parse(self, filename: str, content: bytes) -> tuple[str, dict]:
        suffix = Path(filename).suffix.lower()
        if suffix not in SUPPORTED:
            raise AppError(
                "UNSUPPORTED_DOCUMENT", "仅支持 PDF、DOCX、Markdown 和 TXT", 415
            )
        if not content:
            raise AppError("EMPTY_DOCUMENT", "文件内容为空", 422)
        if len(content) > MAX_DOCUMENT_BYTES:
            raise AppError("DOCUMENT_TOO_LARGE", "文件不能超过 50MB", 413)

        try:
            if suffix == ".pdf":
                reader = PdfReader(BytesIO(content))
                pages = [page.extract_text() or "" for page in reader.pages]
                body = "\n\n".join(part.strip() for part in pages if part.strip())
                metadata = {"pages": len(reader.pages)}
            elif suffix == ".docx":
                document = Document(BytesIO(content))
                paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
                table_rows = [
                    " | ".join(cell.text.strip() for cell in row.cells)
                    for table in document.tables
                    for row in table.rows
                ]
                body = "\n".join([*paragraphs, *table_rows])
                metadata = {"paragraphs": len(document.paragraphs), "tables": len(document.tables)}
            else:
                body = self._decode_text(content)
                metadata = {"characters": len(body)}
        except AppError:
            raise
        except Exception as exc:
            raise AppError("DOCUMENT_PARSE_FAILED", "文档解析失败", 422) from exc

        if not body.strip():
            raise AppError("DOCUMENT_HAS_NO_TEXT", "文档中没有可索引文本", 422)
        return body, metadata

    @staticmethod
    def _decode_text(content: bytes) -> str:
        for encoding in ("utf-8-sig", "utf-8", "gb18030"):
            try:
                return content.decode(encoding)
            except UnicodeDecodeError:
                continue
        raise AppError("DOCUMENT_ENCODING", "文本编码无法识别", 422)

    def import_bytes(self, filename: str, media_type: str, content: bytes) -> tuple[dict, bool]:
        body, metadata = self.parse(filename, content)
        course_id = self._course_id()
        digest = hashlib.sha256(content).hexdigest()
        with self.database.connect() as connection:
            existing = connection.execute(
                "SELECT * FROM documents WHERE course_id = ? AND sha256 = ?",
                (course_id, digest),
            ).fetchone()
        if existing:
            return self._serialize(existing), True

        self.documents_dir.mkdir(parents=True, exist_ok=True)
        suffix = Path(filename).suffix.lower()
        stored_name = f"{digest[:16]}{suffix}"
        stored = self.documents_dir / stored_name
        temporary = stored.with_suffix(stored.suffix + ".tmp")
        temporary.write_bytes(content)
        temporary.replace(stored)
        relative_path = stored.relative_to(self.data_dir).as_posix()

        try:
            with self.database.connect() as connection:
                cursor = connection.execute(
                    """INSERT INTO documents(
                        course_id, title, filename, stored_path, media_type, sha256, body, metadata_json
                    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?)""",
                    (
                        course_id,
                        Path(filename).stem,
                        filename,
                        relative_path,
                        media_type or "application/octet-stream",
                        digest,
                        body,
                        json.dumps(metadata, ensure_ascii=False),
                    ),
                )
                row = connection.execute(
                    "SELECT * FROM documents WHERE id = ?", (cursor.lastrowid,)
                ).fetchone()
        except Exception:
            stored.unlink(missing_ok=True)
            raise
        return self._serialize(row), False

    def list_documents(self) -> list[dict]:
        with self.database.connect() as connection:
            rows = connection.execute(
                "SELECT * FROM documents WHERE course_id = ? ORDER BY id DESC", (self._course_id(),)
            ).fetchall()
        return [self._serialize(row) for row in rows]

    def get_document(self, document_id: int) -> dict:
        with self.database.connect() as connection:
            row = connection.execute(
                "SELECT * FROM documents WHERE id = ? AND course_id = ?",
                (document_id, self._course_id()),
            ).fetchone()
        if not row:
            raise AppError("DOCUMENT_NOT_FOUND", "文档不存在", 404)
        return self._serialize(row)

    def add_highlight(self, document_id: int, values: dict) -> dict:
        self.get_document(document_id)
        with self.database.connect() as connection:
            cursor = connection.execute(
                """INSERT INTO document_highlights(
                    document_id, quote, note, start_offset, end_offset
                ) VALUES (?, ?, ?, ?, ?)""",
                (
                    document_id,
                    values["quote"],
                    values.get("note", ""),
                    values.get("start_offset"),
                    values.get("end_offset"),
                ),
            )
            row = connection.execute(
                "SELECT * FROM document_highlights WHERE id = ?", (cursor.lastrowid,)
            ).fetchone()
        return as_dict(row)

    def search(self, query: str, limit: int = 20) -> list[dict]:
        terms = re.findall(r"[\w\u3400-\u9fff]+", query, flags=re.UNICODE)
        if not terms:
            return []
        fts_query = " AND ".join(f'"{term}"' for term in terms[:10])
        with self.database.connect() as connection:
            rows = connection.execute(
                """SELECT d.id AS document_id, d.title, d.filename,
                    snippet(document_fts, 1, '<mark>', '</mark>', '…', 24) AS snippet,
                    bm25(document_fts) AS score
                FROM document_fts JOIN documents d ON d.id = document_fts.rowid
                WHERE document_fts MATCH ? AND d.course_id = ?
                ORDER BY score LIMIT ?""",
                (fts_query, self._course_id(), limit),
            ).fetchall()
        return [as_dict(row) for row in rows]

    @staticmethod
    def _serialize(row) -> dict:
        item = as_dict(row)
        item["metadata"] = json.loads(item.pop("metadata_json"))
        return item

